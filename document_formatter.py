from io import BytesIO

from docx import Document
from docx.shared import Pt

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from reportlab.lib.units import inch


def format_docx(text: str, document_type: str = "Legal Document") -> bytes:
    """
    Create a DOCX document and return it as bytes.
    """

    document = Document()

    # Title
    title = document.add_paragraph()
    title.alignment = 1

    run = title.add_run(document_type)
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph()

    # Document content
    for line in text.splitlines():

        line = line.strip()

        if line:
            paragraph = document.add_paragraph()
            paragraph.add_run(line)
        else:
            document.add_paragraph()

    # Save into memory
    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


def format_pdf(text: str, document_type: str = "Legal Document") -> bytes:
    """
    Create a PDF document and return it as bytes.
    """

    output = BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.alignment = TA_LEFT

    body_style = styles["BodyText"]
    body_style.leading = 16

    story = []

    # Title
    story.append(
        Paragraph(
            document_type,
            title_style
        )
    )

    story.append(
        Spacer(1, 20)
    )

    # Content
    for line in text.splitlines():

        line = line.strip()

        if line:
            # Escape special HTML characters
            line = (
                line
                .replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )

            story.append(
                Paragraph(
                    line,
                    body_style
                )
            )

            story.append(
                Spacer(1, 8)
            )

        else:
            story.append(
                Spacer(1, 8)
            )

    pdf.build(story)

    output.seek(0)

    return output.getvalue()